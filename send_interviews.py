#!/usr/bin/env python3
"""
Send personalised interview-request emails from a spreadsheet + a text template.

DRY RUN is the default: nothing is sent, previews are printed and saved as .eml
files. Add --send to actually send (you will be asked to confirm).

Data can come from a Google Sheet (read-only, via a service account) or an .xlsx file.
Everyday settings (sheet URL, Gmail address, your name) can live in config.txt.
No AI or agent is involved when this runs. It reads your spreadsheet, fills in
the template and talks to your mail server, and that is all it does.
"""
import argparse
import csv
import getpass
import html
import os
import re
import smtplib
import ssl
import sys
import time
from datetime import datetime, timezone
from email.message import EmailMessage
from email.utils import formataddr, formatdate, make_msgid
from pathlib import Path

# ---- Spreadsheet layout (tab names are matched case-insensitively) ----------
TAB_PUBLICATIONS = "publications"  # publication_id, publication_name, publication_description
TAB_ARTICLES = "articles"          # publication_id, article_id, article_description
TAB_SOURCES = "sources"            # publication_id, article_id, source_name, source_email,
                                   # timeframe, reason_for_the_request

# [UPPERCASE PLACEHOLDERS] in the template map to column names:
# [PUBLICATION NAME] -> publication_name, [REASON FOR THE REQUEST] -> reason_for_the_request
PLACEHOLDER_RE = re.compile(r"\[([A-Z][A-Z0-9 _]*)\]")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)|(https?://[^\s<>\"]+)")

# Fields that may be left blank. If empty, the placeholder is dropped from the email, together with a comma
# and spaces just before it, so "[PUBLICATION NAME], [PUBLICATION DESCRIPTION]." becomes "The Example Review."
OPTIONAL_FIELDS = {"publication_description"}

DEFAULT_SUBJECT = "Interview request for [PUBLICATION NAME]"


# ---- Settings file (config.txt) ----------------------------------------------
CONFIG = {}
NEVER_IN_CONFIG = {"SMTP_PASSWORD"}


def load_config(path, required):
    """Read NAME=value lines from a settings file into CONFIG. Passwords are deliberately ignored."""
    candidates = [Path(path)]
    if not Path(path).is_absolute():
        candidates.append(Path(__file__).resolve().parent / path)
    found = next((c for c in candidates if c.is_file()), None)
    if not found:
        if required:
            die(f"Settings file not found: {path}")
        return None
    for n, line in enumerate(found.read_text(encoding="utf-8-sig").splitlines(), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if "=" not in line:
            print(f"Warning: {found.name} line {n} ignored (expected NAME=value): {line}", file=sys.stderr)
            continue
        key, _, val = line.partition("=")
        key, val = key.strip().upper(), val.strip()
        if len(val) >= 2 and val[0] == val[-1] and val[0] in "\"'":
            val = val[1:-1]
        if key in NEVER_IN_CONFIG:
            print(f"Warning: {key} in {found.name} is ignored for safety. The script asks for it when sending.",
                  file=sys.stderr)
            continue
        CONFIG[key] = val
    return found


def setting(name, default=""):
    """Look a setting up in: terminal environment variable > config.txt > default."""
    if name in os.environ:
        return os.environ[name]
    return CONFIG.get(name, default)


def norm(name):
    return re.sub(r"[^a-z0-9]+", "_", str(name).lower()).strip("_")


def die(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


# ---- Reading the data --------------------------------------------------------
# Each loader returns {tab_name_lowercase: [(row_number, [cell text, ...]), ...]}
def with_link(text, link):
    """A cell with a hyperlink attached (display text -> URL) becomes [text](url)."""
    if link and link not in text:
        return f"[{text}]({link})" if text else link
    return text


def xlsx_cell_text(cell):
    v = cell.value
    if isinstance(v, datetime):
        text = f"{v.day} {v:%B %Y}"
    elif isinstance(v, float) and v.is_integer():
        text = str(int(v))
    elif v is None:
        text = ""
    else:
        text = str(v).strip()
    link = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else None
    return with_link(text, link)


def open_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        die("Reading .xlsx needs openpyxl:  pip install openpyxl")
    if not Path(path).exists():
        die(f"Workbook not found: {path}")
    wb = load_workbook(path)  # not read-only: needed to see hyperlinks
    return {ws.title.lower().strip(): [(r[0].row, [xlsx_cell_text(c) for c in r]) for r in ws.iter_rows()]
            for ws in wb.worksheets}


def open_gsheet(sheet, cred_path):
    """Read every tab of a Google Sheet with a read-only service account."""
    try:
        from google.oauth2 import service_account
        from google.auth.transport.requests import AuthorizedSession
    except ImportError:
        die("Reading Google Sheets needs:  pip install google-auth requests")
    m = re.search(r"/d/([A-Za-z0-9_-]+)", sheet)
    sheet_id = m.group(1) if m else sheet
    if not Path(cred_path).exists():
        die(f"Service-account key file not found: {cred_path} (see README, 'Google Sheets')")
    creds = service_account.Credentials.from_service_account_file(
        cred_path, scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
    resp = AuthorizedSession(creds).get(
        f"https://sheets.googleapis.com/v4/spreadsheets/{sheet_id}",
        params={"includeGridData": "true",
                "fields": "sheets(properties(title),data(startRow,rowData(values(formattedValue,hyperlink))))"},
        timeout=60)
    if resp.status_code != 200:
        hint = ""
        if resp.status_code in (403, 404):
            hint = (f"\nShare the sheet (Viewer) with the service account: {creds.service_account_email}"
                    "\nAlso check the Google Sheets API is enabled for that Google Cloud project.")
        die(f"Google Sheets API returned {resp.status_code}: {resp.text[:300]}{hint}")
    tabs = {}
    for sh in resp.json().get("sheets", []):
        rows = []
        for data in sh.get("data", []):
            start = data.get("startRow", 0)
            for i, rd in enumerate(data.get("rowData", [])):
                cells = [with_link(c.get("formattedValue", "").strip(), c.get("hyperlink"))
                         for c in rd.get("values", [])]
                rows.append((start + i + 1, cells))
        tabs[sh["properties"]["title"].lower().strip()] = rows
    return tabs


def read_sheet(tabs, tab):
    if tab not in tabs:
        die(f"No '{tab}' tab found (found: {', '.join(tabs) or 'none'}).")
    rows = tabs[tab]
    if not rows:
        return []
    headers = [norm(h) for h in rows[0][1]]
    out = []
    for rownum, vals in rows[1:]:
        data = {h: v for h, v in zip(headers, vals) if h}
        if any(data.values()):
            out.append((rownum, data))
    return out


def index_rows(rows, keys, tab, problems):
    idx = {}
    for rownum, data in rows:
        key = tuple(data.get(k, "") for k in keys)
        if not all(key):
            problems.append(f"{tab} row {rownum}: missing {' / '.join(keys)}")
            continue
        if key in idx:
            problems.append(f"{tab} row {rownum}: duplicate {' / '.join(keys)} = {' / '.join(key)}")
            continue
        idx[key] = data
    return idx


# ---- Template and link handling ---------------------------------------------
def load_template(path):
    text = Path(path).read_text(encoding="utf-8")
    lines = text.splitlines()
    if lines and lines[0].lower().startswith("subject:"):
        subject = lines[0][len("subject:"):].strip()
        body = "\n".join(lines[1:]).strip("\n")
    else:
        subject, body = DEFAULT_SUBJECT, text.strip("\n")
    return subject, body


def drop_empty_optional(template, ctx):
    """Remove empty optional placeholders plus a comma/spaces directly before them (never across a line break)."""
    removed = []

    def sub(m):
        name = norm(m.group(1)[1:-1])
        if name in OPTIONAL_FIELDS and not ctx.get(name):
            removed.append(name)
            return ""
        return m.group(0)

    out = re.sub(r"(?:[ \t]*,)?[ \t]*(\[[A-Z][A-Z0-9 _]*\])", sub, template)
    if removed:
        out = re.sub(r"\n{3,}", "\n\n", out)  # don't leave a gap if the placeholder was on its own line
    return out


TPL = "TEMPLATE|"  # marks a template-structure problem in the `missing` set
TAG_RE = re.compile(r"(\{IF [A-Z][A-Z0-9 _]*\}|\{ELSE\}|\{END IF\})")
FALSE_WORDS = {"", "no", "n", "false", "0", "none", "n/a"}  # anything else counts as "yes"


def resolve_conditionals(template, ctx, missing, notes):
    """
    {IF FIELD} text {END IF}             keep the text only if FIELD has a value
    {IF FIELD} a {ELSE} b {END IF}       use a if FIELD has a value, otherwise b
    A field counts as "no" if it is blank or says no / n / false / 0 / none / n/a (any capitals).
    Blocks can be nested. Only the branch that is used is checked and reported.
    """
    tokens = TAG_RE.split(template)  # even positions: plain text, odd positions: {IF ..} / {ELSE} / {END IF}
    pos = 0
    dropped = False

    def parse_seq(active):
        """Read text and nested blocks up to the next {ELSE} / {END IF} (left for the caller)."""
        nonlocal pos
        out = []
        while pos < len(tokens):
            tok = tokens[pos]
            if pos % 2 == 0:
                if active:
                    out.append(tok)
                pos += 1
            elif tok in ("{ELSE}", "{END IF}"):
                break
            else:
                out.append(parse_block(tok, active))
        return "".join(out)

    def parse_block(tag, active):
        nonlocal pos, dropped
        label = tag[4:-1]
        name = norm(label)
        pos += 1
        known_field = name in ctx or name in OPTIONAL_FIELDS
        if active and not known_field:
            missing.add(TPL + f"{{IF {label}}} (no column with that name)")
        truthy = known_field and ctx.get(name, "").strip().lower() not in FALSE_WORDS
        if_text = parse_seq(active and truthy)
        else_text, has_else = "", False
        if pos < len(tokens) and tokens[pos] == "{ELSE}":
            has_else = True
            pos += 1
            else_text = parse_seq(active and known_field and not truthy)
            while pos < len(tokens) and tokens[pos] == "{ELSE}":
                missing.add(TPL + f"{{IF {label}}} has more than one {{ELSE}}")
                pos += 1
                parse_seq(False)
        if pos < len(tokens) and tokens[pos] == "{END IF}":
            pos += 1
        else:
            missing.add(TPL + f"{{IF {label}}} has no {{END IF}}")
        if not (active and known_field):
            return ""
        if truthy:
            return if_text
        dropped = True
        if has_else:
            notes.add(f"{{IF {label}}}: {name} is blank or no, so the {{ELSE}} text was used")
        else:
            notes.add(f"left out the text inside {{IF {label}}} because {name} is blank or no")
        return else_text

    out = []
    while pos < len(tokens):
        out.append(parse_seq(True))
        if pos < len(tokens):
            missing.add(TPL + "a stray {ELSE} or {END IF} with no matching {IF FIELD NAME}")
            pos += 1
    result = "".join(out)
    if re.search(r"\{\s*(?:IF|ELSE|END\s*IF)", result, re.I):
        missing.add(TPL + "{IF ...} / {ELSE} / {END IF} written wrongly (use capitals, e.g. {IF FIELD NAME} ... {END IF})")
    if dropped:
        result = re.sub(r"\n{3,}", "\n\n", result)  # don't leave a gap if a whole paragraph was removed
    return result


def render(template, ctx, missing, notes=None):
    notes = notes if notes is not None else set()
    template = resolve_conditionals(template, ctx, missing, notes)
    template = drop_empty_optional(template, ctx)

    def sub(m):
        val = ctx.get(norm(m.group(1)), "")
        if not val:
            missing.add(m.group(1))
            return m.group(0)
        # Avoid ".." when the template puts a full stop right after a value ending in one
        if val.endswith(".") and m.string[m.end():m.end() + 1] == ".":
            val = val[:-1]
        return val
    return PLACEHOLDER_RE.sub(sub, template)


def to_plain(text):
    def sub(m):
        if m.group(3):
            return m.group(3)
        label, url = m.group(1), m.group(2)
        return url if label.strip() == url else f"{label} ({url})"
    return LINK_RE.sub(sub, text)


def to_subject(text):
    """Subject lines carry no hyperlinks: [text](url) keeps just the text, and bare URLs are dropped."""
    text = LINK_RE.sub(lambda m: m.group(1) or "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return re.sub(r"\s+([,.;:])", r"\1", text)


def to_html(text):
    out, pos = [], 0
    for m in LINK_RE.finditer(text):
        out.append(html.escape(text[pos:m.start()]))
        if m.group(3):
            url = m.group(3).rstrip(".,;:!?)")
            trailing = m.group(3)[len(url):]
            out.append(f'<a href="{html.escape(url)}">{html.escape(url)}</a>{html.escape(trailing)}')
        else:
            out.append(f'<a href="{html.escape(m.group(2))}">{html.escape(m.group(1))}</a>')
        pos = m.end()
    out.append(html.escape(text[pos:]))
    escaped = "".join(out)
    paras = [p.replace("\n", "<br>\n") for p in re.split(r"\n\s*\n", escaped.strip())]
    body = "\n".join(f"<p>{p}</p>" for p in paras)
    return f'<div style="font-family: Arial, Helvetica, sans-serif; font-size: 14px;">\n{body}\n</div>'


# ---- Building the emails -----------------------------------------------------
def build_messages(args, cfg):
    tabs = open_gsheet(args.sheet, args.credentials) if args.sheet else open_xlsx(args.workbook)
    problems = []
    pubs = index_rows(read_sheet(tabs, TAB_PUBLICATIONS), ("publication_id",), TAB_PUBLICATIONS, problems)
    arts = index_rows(read_sheet(tabs, TAB_ARTICLES), ("publication_id", "article_id"), TAB_ARTICLES, problems)
    sources = read_sheet(tabs, TAB_SOURCES)
    subject_tpl, body_tpl = load_template(args.template)

    sent = load_log(args.log)
    items = []
    for rownum, src in sources:
        where = f"{TAB_SOURCES} row {rownum}"
        pid, aid = src.get("publication_id", ""), src.get("article_id", "")
        email = src.get("source_email", "")
        if args.only_publication and pid != args.only_publication:
            continue
        if args.only_article and aid != args.only_article:
            continue
        if not EMAIL_RE.match(email):
            problems.append(f"{where}: invalid or missing source_email '{email}'")
            continue
        if (pid, aid, email.lower()) in sent and not args.resend:
            items.append(("skip", f"{email} ({pid}/{aid}) already in {args.log}"))
            continue
        pub, art = pubs.get((pid,)), arts.get((pid, aid))
        if pub is None:
            problems.append(f"{where}: publication_id '{pid}' not found in {TAB_PUBLICATIONS}")
            continue
        if art is None:
            problems.append(f"{where}: publication_id/article_id '{pid}/{aid}' not found in {TAB_ARTICLES}")
            continue

        ctx = {**pub, **art, **src}
        missing, notes = set(), set()
        # A subject_line (or subject) column in the spreadsheet wins; otherwise the template's Subject: line
        subject = render(ctx.get("subject_line") or ctx.get("subject") or subject_tpl, ctx, missing, notes)
        body = render(body_tpl, ctx, missing, notes)
        if missing:
            fields = sorted(f"[{m}]" for m in missing if not m.startswith(TPL))
            for issue in sorted(m[len(TPL):] for m in missing if m.startswith(TPL)):
                msg = f"template problem: {issue}"  # same for every row, so report it once
                if msg not in problems:
                    problems.append(msg)
            if fields:
                problems.append(f"{where}: empty or unknown value for {', '.join(fields)}")
            continue

        subject = to_subject(subject)
        msg = EmailMessage()
        msg["From"] = formataddr((cfg["from_name"], cfg["from_addr"]))
        msg["To"] = formataddr((src.get("source_name", ""), email))
        msg["Subject"] = subject
        if cfg["reply_to"]:
            msg["Reply-To"] = cfg["reply_to"]
        msg["Date"] = formatdate(localtime=True)
        msg["Message-ID"] = make_msgid(domain=cfg["from_addr"].split("@")[-1])
        msg.set_content(to_plain(body))
        msg.add_alternative(to_html(body), subtype="html")
        items.append(("send", {"msg": msg, "pid": pid, "aid": aid, "email": email,
                               "subject": subject, "plain": to_plain(body), "notes": sorted(notes)}))
    return items, problems


# ---- Sent log (so re-running never emails the same person twice) -------------
def load_log(path):
    done = set()
    if Path(path).exists():
        with open(path, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                done.add((row["publication_id"], row["article_id"], row["source_email"].lower()))
    return done


def append_log(path, item):
    new = not Path(path).exists()
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["sent_at", "publication_id", "article_id", "source_email", "subject"])
        w.writerow([datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    item["pid"], item["aid"], item["email"], item["subject"]])


# ---- SMTP --------------------------------------------------------------------
def connect(cfg):
    ctx = ssl.create_default_context()
    if cfg["loopback_insecure"]:
        # Proton Bridge uses a self-signed certificate on 127.0.0.1 only.
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
    if cfg["security"] == "ssl":
        server = smtplib.SMTP_SSL(cfg["host"], cfg["port"], context=ctx, timeout=60)
    else:
        server = smtplib.SMTP(cfg["host"], cfg["port"], timeout=60)
        server.starttls(context=ctx)
    server.login(cfg["user"], cfg["password"])
    return server


def get_config(args):
    if args.bridge:
        host, port, insecure = "127.0.0.1", int(setting("SMTP_PORT", "1025")), True
    elif args.gmail:
        host, port, insecure = "smtp.gmail.com", 587, False
    else:
        host = setting("SMTP_HOST", "smtp.protonmail.ch")
        port = int(setting("SMTP_PORT", "587"))
        insecure = False
    user = setting("SMTP_USER", "").strip()
    return {
        "host": host, "port": port, "user": user,
        "password": os.environ.get("SMTP_PASSWORD", ""),
        "security": setting("SMTP_SECURITY", "starttls").lower(),
        "loopback_insecure": insecure,
        "from_addr": setting("FROM_ADDRESS").strip() or user,
        "from_name": setting("FROM_NAME", "Kari"),
        "reply_to": setting("REPLY_TO", "kari.mcmahon.freelance@proton.me").strip(),
    }


def main():
    pre = argparse.ArgumentParser(add_help=False)
    pre.add_argument("--config")
    known, _ = pre.parse_known_args()
    config_file = load_config(known.config or "config.txt", required=bool(known.config))

    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--config", help="settings file (default config.txt)")
    p.add_argument("--sheet", default=setting("SHEET_URL") or setting("GOOGLE_SHEET_ID") or None,
                   help="Google Sheet URL or ID (SHEET_URL in config.txt); takes priority over --workbook")
    p.add_argument("--credentials",
                   default=setting("CREDENTIALS") or setting("GOOGLE_APPLICATION_CREDENTIALS") or "service-account.json",
                   help="service-account JSON key file (default service-account.json)")
    p.add_argument("--workbook", default="interviews.xlsx", help=".xlsx file, used when no --sheet is given")
    p.add_argument("--template", default="template.txt")
    p.add_argument("--send", action="store_true", help="actually send (default is a dry run)")
    p.add_argument("--bridge", action="store_true", help="use Proton Mail Bridge on 127.0.0.1:1025")
    p.add_argument("--gmail", action="store_true", help="use smtp.gmail.com:587 (needs a Google app password)")
    p.add_argument("--only-publication", metavar="ID")
    p.add_argument("--only-article", metavar="ID")
    p.add_argument("--resend", action="store_true", help="ignore the sent log")
    p.add_argument("--delay", type=int, default=15, help="seconds between emails (default 15)")
    p.add_argument("--log", default="sent_log.csv")
    p.add_argument("--preview-dir", default="preview")
    args = p.parse_args()

    if config_file:
        print(f"Using settings from {config_file}")
    cfg = get_config(args)
    if not cfg["from_addr"]:
        if args.send:
            die("Set SMTP_USER (your sending address) in config.txt.")
        cfg["from_addr"] = "you@example.com"  # placeholder, dry run only

    items, problems = build_messages(args, cfg)
    if problems:
        print("Fix these problems in the spreadsheet or template first (nothing was sent):\n")
        for pr in problems:
            print(f"  - {pr}")
        sys.exit(1)

    to_send = [i[1] for i in items if i[0] == "send"]
    for kind, info in items:
        if kind == "skip":
            print(f"Skipping: {info}")

    Path(args.preview_dir).mkdir(exist_ok=True)
    for n, it in enumerate(to_send, 1):
        print(f"\n{'=' * 70}\n[{n}/{len(to_send)}] To: {it['msg']['To']}\nSubject: {it['subject']}")
        for note in it["notes"]:
            print(f"Note: {note}")
        print("-" * 70)
        print(it["plain"])
        (Path(args.preview_dir) / f"{n:02d}_{it['email']}.eml").write_bytes(bytes(it["msg"]))

    if not to_send:
        print("\nNothing to send.")
        return
    if not args.send:
        print(f"\nDRY RUN: {len(to_send)} email(s) previewed, none sent. "
              f"Previews saved in ./{args.preview_dir}/. Re-run with --send to send.")
        return

    print(f"\nAbout to send {len(to_send)} email(s) from {cfg['from_addr']}:")
    for it in to_send:
        print(f"  {it['email']}")
    if input("\nType SEND to continue: ").strip() != "SEND":
        print("Cancelled.")
        return

    if not cfg["user"]:
        die("Set SMTP_USER (your sending address) in config.txt.")
    if not cfg["password"]:
        if args.gmail:
            label = "Gmail APP PASSWORD (16 characters from myaccount.google.com/apppasswords, not your normal password)"
        elif args.bridge:
            label = "Proton Bridge password (from the Bridge app, not your Proton password)"
        else:
            label = "SMTP password / token"
        cfg["password"] = getpass.getpass(f"{label}\nInput is hidden: ")

    server = connect(cfg)
    ok = failed = 0
    for n, it in enumerate(to_send, 1):
        try:
            try:
                server.send_message(it["msg"])
            except smtplib.SMTPServerDisconnected:
                server = connect(cfg)
                server.send_message(it["msg"])
            append_log(args.log, it)
            ok += 1
            print(f"[{n}/{len(to_send)}] sent to {it['email']}")
        except Exception as e:  # keep going; failures are not logged so a re-run retries them
            failed += 1
            print(f"[{n}/{len(to_send)}] FAILED for {it['email']}: {e}", file=sys.stderr)
        if n < len(to_send):
            time.sleep(args.delay)
    try:
        server.quit()
    except Exception:
        pass
    print(f"\nDone. Sent: {ok}, failed: {failed}. Log: {args.log}")


if __name__ == "__main__":
    main()
